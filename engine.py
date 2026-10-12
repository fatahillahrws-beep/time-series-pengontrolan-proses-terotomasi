"""
engine.py — Mesin komputasi (tanpa Streamlit)
Perbandingan 9 metode prediksi multivariat NPK Phonska IV + control chart MEWMA & Adaptive-DHWMA.
Dikonversi dari notebook `NPK_Perbandingan_9_Metode.ipynb`.
"""
import os, io, time, random, itertools, warnings, gc, json, hashlib, pickle
import numpy as np
import pandas as pd

warnings.filterwarnings('ignore')

METHOD_ORDER = ['VAR', 'SVR', 'Random Forest', 'XGBoost', 'LightGBM', 'LSTM', 'RNN', 'GRU', 'TCN']
KATEGORI = {'VAR': 'Statistik', 'SVR': 'Machine Learning', 'Random Forest': 'Machine Learning',
            'XGBoost': 'Machine Learning', 'LightGBM': 'Machine Learning',
            'LSTM': 'Deep Learning', 'RNN': 'Deep Learning', 'GRU': 'Deep Learning', 'TCN': 'Deep Learning'}

XGB_PARAMS = {
    'H2O':  dict(colsample_bytree=0.643157, learning_rate=0.010987, max_depth=3, min_child_weight=4, n_estimators=295, reg_lambda=2.46055,  subsample=0.655733),
    'N':    dict(colsample_bytree=0.739466, learning_rate=0.013339, max_depth=3, min_child_weight=1, n_estimators=750, reg_lambda=7.472398, subsample=0.655008),
    'P2O5': dict(colsample_bytree=0.720351, learning_rate=0.023474, max_depth=3, min_child_weight=1, n_estimators=201, reg_lambda=1.012414, subsample=0.620592),
    'K2O':  dict(colsample_bytree=0.643157, learning_rate=0.010987, max_depth=3, min_child_weight=4, n_estimators=295, reg_lambda=2.46055,  subsample=0.655733),
}
LGBM_PARAMS = {
    'H2O':  dict(colsample_bytree=0.975421, learning_rate=0.010023, max_depth=6, n_estimators=360, reg_lambda=1.67227,  subsample=0.602827),
    'N':    dict(colsample_bytree=0.859076, learning_rate=0.010016, max_depth=7, n_estimators=458, reg_lambda=0.213458, subsample=0.813636),
    'P2O5': dict(colsample_bytree=0.975421, learning_rate=0.010023, max_depth=6, n_estimators=360, reg_lambda=1.67227,  subsample=0.602827),
    'K2O':  dict(colsample_bytree=0.975421, learning_rate=0.010023, max_depth=6, n_estimators=360, reg_lambda=1.67227,  subsample=0.602827),
}
CACHE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.cache_hasil')


def set_seed(s, tf=None):
    os.environ['PYTHONHASHSEED'] = str(s); random.seed(s); np.random.seed(s)
    if tf is not None:
        tf.random.set_seed(s)


# ----------------------------------------------------------------------------- konteks
class Ctx:
    def __init__(self, df, cfg, progress):
        self.cfg, self.progress = cfg, progress
        self.vars = list(cfg['vars'])
        self.DATA = df[self.vars].values.astype(float)
        self.N = len(self.DATA)
        self.N_TRAIN = int(cfg['train_ratio'] * self.N)
        self.Y_TEST = self.DATA[self.N_TRAIN:]
        self.T_TEST = pd.to_datetime(df[cfg['time_col']].iloc[self.N_TRAIN:]).values
        self.SD_TRAIN = self.DATA[:self.N_TRAIN].std(0, ddof=1)
        self.results, self.preds, self.timing, self.extra, self.warnings = [], {}, {}, {}, []
        self.base, self.span = 0.0, 1.0

    def tick(self, msg, frac=0.0):
        self.progress(msg, min(0.99, self.base + self.span * frac))

    def make_windows(self, arr, w):
        X = np.stack([arr[i - w:i] for i in range(w, len(arr))])
        return X, arr[w:], np.arange(w, len(arr))

    def split_windows(self, arr, w):
        X, y, idx = self.make_windows(arr, w)
        tr = idx < self.N_TRAIN
        return X[tr], y[tr], X[~tr], y[~tr]

    def evaluate(self, y_true, y_pred):
        out = {}
        for j, v in enumerate(self.vars):
            e = y_true[:, j] - y_pred[:, j]
            out[f'RMSE_{v}'] = float(np.sqrt(np.mean(e ** 2)))
            out[f'MAPE_{v}'] = float(np.mean(np.abs(e / y_true[:, j])) * 100)
        out['RMSE_avg'] = float(np.mean([out[f'RMSE_{v}'] for v in self.vars]))
        out['MAPE_avg'] = float(np.mean([out[f'MAPE_{v}'] for v in self.vars]))
        out['RMSE_norm'] = float(np.mean([out[f'RMSE_{v}'] / self.SD_TRAIN[j] for j, v in enumerate(self.vars)]))
        return out

    def record(self, method, w, config, y_pred, note=''):
        assert y_pred.shape == self.Y_TEST.shape
        m = self.evaluate(self.Y_TEST, y_pred)
        self.results.append({'Metode': method, 'Kategori': KATEGORI[method], 'Window': int(w),
                             'Konfigurasi': config, 'Catatan': note, **m})
        self.preds[(method, int(w))] = y_pred


# ----------------------------------------------------------------------------- VAR
def run_var(ctx):
    from statsmodels.tsa.api import VAR
    from statsmodels.tsa.stattools import adfuller
    cfg, V, NT = ctx.cfg, ctx.vars, ctx.N_TRAIN
    train_df = pd.DataFrame(ctx.DATA[:NT], columns=V)
    k_adf = int(np.trunc((NT - 1) ** (1 / 3)))
    rows = []
    for v in V:
        stat, pval = adfuller(train_df[v], maxlag=k_adf, regression='ct', autolag=None)[:2]
        rows.append({'Variabel': v, 'ADF': stat, 'P_Value': pval,
                     'Keputusan': 'Stasioner' if pval < 0.05 else 'Tidak Stasioner'})
    ctx.extra['VAR_ADF'] = pd.DataFrame(rows)

    maxlags = int(max(2, min(cfg['var_maxlags'], NT // (2 * (len(V) + 1)))))
    sel = VAR(train_df).select_order(maxlags=maxlags, trend='c')
    p_aic, p_bic = max(1, int(sel.selected_orders['aic'])), max(1, int(sel.selected_orders['bic']))
    ctx.extra['P_AIC'], ctx.extra['P_BIC'] = p_aic, p_bic
    ctx.extra['VAR_LAG'] = pd.DataFrame({'Lag': range(1, maxlags + 1),
                                         'AIC': np.array(sel.ics['aic'])[1:], 'BIC': np.array(sel.ics['bic'])[1:]})
    stab, plist = [], sorted(set(list(cfg['windows']) + [p_aic]))
    for i, p in enumerate(plist):
        ctx.tick(f'VAR lag {p}', i / len(plist))
        fit = VAR(train_df).fit(p, trend='c')
        max_mod = float(np.max(np.abs(1 / fit.roots))) if len(fit.roots) else 0.0
        pred = np.vstack([fit.forecast(ctx.DATA[ctx.N_TRAIN + k - p:ctx.N_TRAIN + k], 1)[0]
                          for k in range(ctx.N - ctx.N_TRAIN)])
        stab.append({'Lag': p, 'Modulus_akar_maks': max_mod, 'Stabil': bool(max_mod < 1)})
        ctx.record('VAR', p, f'VAR({p}), konstanta', pred, 'lag optimal AIC (seperti VAR.R)' if p == p_aic else '')
    ctx.extra['VAR_STAB'] = pd.DataFrame(stab)


# ----------------------------------------------------------------------------- SVR
def run_svr(ctx):
    from sklearn.svm import SVR
    cfg, V = ctx.cfg, ctx.vars
    grid = list(itertools.product(cfg['svr_C'], cfg['svr_gamma'], cfg['svr_eps']))

    def fit_scaler(X):
        mu, sd = X.mean(0), X.std(0, ddof=1)
        sd[~np.isfinite(sd) | (sd == 0)] = 1
        return mu, sd

    def folds_(n, k=3, prop=0.15):
        h = max(10, int(np.floor(n * prop))); start = n - k * h
        return [(np.arange(0, start + i * h), np.arange(start + i * h, start + (i + 1) * h)) for i in range(k)]

    tuning, best_rows = [], []
    for wi, w in enumerate(cfg['windows']):
        Xtr, ytr, Xte, _ = ctx.split_windows(ctx.DATA, w)
        Xtr, Xte = Xtr.reshape(len(Xtr), -1), Xte.reshape(len(Xte), -1)
        folds = folds_(len(Xtr))
        mu_x, sd_x = fit_scaler(Xtr)
        Xtr_s, Xte_s = (Xtr - mu_x) / sd_x, (Xte - mu_x) / sd_x
        pred = np.zeros((len(Xte), len(V)))
        for j, v in enumerate(V):
            ctx.tick(f'SVR window {w} • {v}', (wi + j / len(V)) / len(cfg['windows']))
            y = ytr[:, j]; rows = []
            for C, g, eps in grid:
                rm = []
                for tr, va in folds:
                    m_, s_ = fit_scaler(Xtr[tr])
                    my, sy = y[tr].mean(), (y[tr].std(ddof=1) or 1)
                    mdl = SVR(kernel='rbf', C=C, gamma=g, epsilon=eps).fit((Xtr[tr] - m_) / s_, (y[tr] - my) / sy)
                    yh = mdl.predict((Xtr[va] - m_) / s_) * sy + my
                    rm.append(np.sqrt(np.mean((y[va] - yh) ** 2)))
                rows.append({'Window': w, 'Variabel': v, 'C': C, 'gamma': g, 'epsilon': eps, 'RMSE_CV': float(np.mean(rm))})
            tun = pd.DataFrame(rows).sort_values('RMSE_CV'); tuning.append(tun)
            b = tun.iloc[0]; best_rows.append(b.to_dict())
            my, sy = y.mean(), (y.std(ddof=1) or 1)
            mdl = SVR(kernel='rbf', C=b.C, gamma=b.gamma, epsilon=b.epsilon).fit(Xtr_s, (y - my) / sy)
            pred[:, j] = mdl.predict(Xte_s) * sy + my
        cfgtxt = '; '.join(f"{r['Variabel']}: C={r['C']:g}, gamma={r['gamma']:g}, eps={r['epsilon']:g}"
                           for r in best_rows if r['Window'] == w)
        ctx.record('SVR', w, cfgtxt, pred)
    ctx.extra['SVR_TUNING'] = pd.concat(tuning, ignore_index=True)
    ctx.extra['SVR_BEST'] = pd.DataFrame(best_rows)


# ----------------------------------------------------------------------------- Random Forest
def run_rf(ctx):
    from sklearn.ensemble import RandomForestRegressor
    from sklearn.model_selection import RandomizedSearchCV, TimeSeriesSplit
    from sklearn.preprocessing import StandardScaler
    cfg = ctx.cfg
    space = {'n_estimators': [100, 200, 300, 500], 'max_depth': [None, 4, 6, 8, 12, 20],
             'min_samples_split': [2, 5, 10, 20], 'min_samples_leaf': [1, 2, 4, 8],
             'max_features': ['sqrt', 0.3, 0.5, 0.75, 1.0]}
    sc = StandardScaler().fit(ctx.DATA[:ctx.N_TRAIN]); DS = sc.transform(ctx.DATA)
    best_rows = []
    for wi, w in enumerate(cfg['windows']):
        ctx.tick(f'Random Forest window {w}', wi / len(cfg['windows']))
        Xtr, ytr, Xte, _ = ctx.split_windows(DS, w)
        Xtr, Xte = Xtr.reshape(len(Xtr), -1), Xte.reshape(len(Xte), -1)
        rs = RandomizedSearchCV(RandomForestRegressor(random_state=cfg['seed'], n_jobs=1), space, n_iter=cfg['rf_iter'],
                                cv=TimeSeriesSplit(n_splits=4), scoring='neg_root_mean_squared_error',
                                random_state=cfg['seed'], n_jobs=-1, refit=False).fit(Xtr, ytr)
        best = rs.best_params_
        mdl = RandomForestRegressor(**best, random_state=cfg['seed'], n_jobs=-1).fit(Xtr, ytr)
        pred = sc.inverse_transform(mdl.predict(Xte))
        best_rows.append({'Window': w, **{k: ('None' if v is None else v) for k, v in best.items()},
                          'CV_RMSE_std': -rs.best_score_})
        ctx.record('Random Forest', w, ', '.join(f'{k}={v}' for k, v in best.items()), pred)
    ctx.extra['RF_BEST'] = pd.DataFrame(best_rows)


# ----------------------------------------------------------------------------- XGBoost / LightGBM
def make_gbm(name, v, seed):
    if name == 'XGBoost':
        from xgboost import XGBRegressor
        return XGBRegressor(**XGB_PARAMS[v], random_state=seed, n_jobs=-1)
    from lightgbm import LGBMRegressor
    return LGBMRegressor(**LGBM_PARAMS[v], subsample_freq=1, random_state=seed, n_jobs=-1, verbose=-1)


def run_gbm(ctx, name):
    cfg = ctx.cfg
    for wi, w in enumerate(cfg['windows']):
        ctx.tick(f'{name} window {w}', wi / len(cfg['windows']))
        Xtr, ytr, Xte, _ = ctx.split_windows(ctx.DATA, w)
        Xtr, Xte = Xtr.reshape(len(Xtr), -1), Xte.reshape(len(Xte), -1)
        pred = np.column_stack([make_gbm(name, v, cfg['seed']).fit(Xtr, ytr[:, j]).predict(Xte)
                                for j, v in enumerate(ctx.vars)])
        ctx.record(name, w, 'hyperparameter tetap per variabel (tanpa tuning)', pred)


def hp_table():
    return pd.concat([pd.DataFrame([{'Model': 'XGBoost', 'Variabel': v, **p} for v, p in XGB_PARAMS.items()]),
                      pd.DataFrame([{'Model': 'LightGBM', 'Variabel': v, **p} for v, p in LGBM_PARAMS.items()])],
                     ignore_index=True)


# ----------------------------------------------------------------------------- Deep Learning
def run_dl(ctx, name):
    import tensorflow as tf
    from tensorflow.keras import layers, models, callbacks, optimizers
    from sklearn.preprocessing import MinMaxScaler
    tf.get_logger().setLevel('ERROR')
    cfg, NF = ctx.cfg, len(ctx.vars)
    mm = MinMaxScaler().fit(ctx.DATA[:ctx.N_TRAIN]); DMM = mm.transform(ctx.DATA)

    def build_rec(cell, w, units, dropout, lr):
        inp = layers.Input(shape=(w, NF))
        x = layers.Dropout(dropout)(cell(units)(inp))
        m = models.Model(inp, layers.Dense(NF)(x)); m.compile(optimizers.Adam(lr), 'mse'); return m

    def tcn_block(x, f, k, d, drop):
        sc = x
        for _ in range(2):
            x = layers.Conv1D(f, k, padding='causal', dilation_rate=d)(x)
            x = layers.LayerNormalization()(x); x = layers.Activation('relu')(x); x = layers.SpatialDropout1D(drop)(x)
        if sc.shape[-1] != f: sc = layers.Conv1D(f, 1, padding='same')(sc)
        return layers.Add()([x, sc])

    def build_tcn(w, units, dropout, lr, k=2):
        nb, rf = 0, 1
        while rf < w: rf += 2 * (k - 1) * 2 ** nb; nb += 1
        inp = layers.Input(shape=(w, NF)); x = inp
        for b in range(max(nb, 1)): x = tcn_block(x, units, k, 2 ** b, dropout)
        x = layers.Lambda(lambda t: t[:, -1, :], output_shape=(units,))(x)
        m = models.Model(inp, layers.Dense(NF)(x)); m.compile(optimizers.Adam(lr), 'mse'); return m

    builders = {'LSTM': lambda w, **c: build_rec(layers.LSTM, w, **c),
                'RNN': lambda w, **c: build_rec(layers.SimpleRNN, w, **c),
                'GRU': lambda w, **c: build_rec(layers.GRU, w, **c),
                'TCN': lambda w, **c: build_tcn(w, **c)}
    grid = [dict(units=int(u), dropout=cfg['dl_dropout'], lr=float(lr)) for u in cfg['dl_units'] for lr in cfg['dl_lrs']]
    total, done = len(cfg['windows']) * len(grid), 0
    rows = ctx.extra.setdefault('DL_RESULTS', [])
    for w in cfg['windows']:
        Xtr, ytr, Xte, _ = ctx.split_windows(DMM, w)
        best = None
        for c in grid:
            ctx.tick(f'{name} window {w} • units={c["units"]}, lr={c["lr"]}', done / total); done += 1
            tf.keras.backend.clear_session(); gc.collect(); set_seed(cfg['seed'], tf)
            mdl = builders[name](w, **c)
            h = mdl.fit(Xtr, ytr, validation_split=0.1, shuffle=False, epochs=cfg['dl_epochs'], batch_size=32, verbose=0,
                        callbacks=[callbacks.EarlyStopping('val_loss', patience=cfg['dl_patience'], restore_best_weights=True),
                                   callbacks.ReduceLROnPlateau('val_loss', factor=0.5, patience=5, min_lr=1e-5)])
            pred = mm.inverse_transform(mdl.predict(Xte, verbose=0))
            m = ctx.evaluate(ctx.Y_TEST, pred)
            rows.append({'Metode': name, 'Window': w, **c, 'batch_size': 32, 'epochs_run': len(h.history['loss']), **m})
            if best is None or m[cfg['select_by']] < best[0][cfg['select_by']]: best = (m, c, pred)
        ctx.record(name, w, f"units={best[1]['units']}, lr={best[1]['lr']}, dropout={cfg['dl_dropout']}, batch=32", best[2])


RUNNERS = {'VAR': run_var, 'SVR': run_svr, 'Random Forest': run_rf,
           'XGBoost': lambda c: run_gbm(c, 'XGBoost'), 'LightGBM': lambda c: run_gbm(c, 'LightGBM'),
           'LSTM': lambda c: run_dl(c, 'LSTM'), 'RNN': lambda c: run_dl(c, 'RNN'),
           'GRU': lambda c: run_dl(c, 'GRU'), 'TCN': lambda c: run_dl(c, 'TCN')}


# ----------------------------------------------------------------------------- Control chart
def mewma_stat(X, lam):
    B, T, p = X.shape; Z = np.zeros((B, p), X.dtype); out = np.empty((B, T), np.float32)
    for t in range(T):
        Z = lam * X[:, t] + (1 - lam) * Z
        out[:, t] = (Z * Z).sum(1) / (lam / (2 - lam) * (1 - (1 - lam) ** (2 * (t + 1))))
    return out


def adhwma_stat(X, lmin, lmax, gamma):
    B, T, p = X.shape; sx = np.zeros((B, p), X.dtype); sh = np.zeros((B, p), X.dtype)
    out = np.empty((B, T), np.float32)
    for t in range(T):
        x = X[:, t]; mx = sx / t if t else 0.0; mh = sh / t if t else 0.0
        lam = (lmin + (lmax - lmin) * (1 - np.exp(-gamma * (x * x).sum(1) / p))).astype(X.dtype)[:, None]
        H = lam * x + (1 - lam) * mx; Q = lam * H + (1 - lam) * mh
        out[:, t] = (Q * Q).sum(1); sx += x; sh += H
    return out


def run_length(stat, h):
    ex = stat > h; sig = ex.any(1)
    return np.where(sig, ex.argmax(1) + 1, stat.shape[1]), sig


def arl_hat(stat, h):
    rl, sig = run_length(stat, h); return float(rl.sum() / max(sig.sum(), 1))


def find_ucl(stat, arl0):
    lo, hi = 0.0, float(stat.max())
    for _ in range(45):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if arl_hat(stat, mid) < arl0 else (lo, mid)
    return (lo + hi) / 2


def standardize(E):
    mu, S = E.mean(0), np.cov(E, rowvar=False)
    L = np.linalg.cholesky(S + 1e-8 * np.eye(E.shape[1]))
    return np.linalg.solve(L, (E - mu).T).T


def run_control_charts(ctx, best_keys):
    cfg = ctx.cfg
    fns = {'MEWMA': lambda X: mewma_stat(X, cfg['lambda_mewma']),
           'Adaptive-DHWMA': lambda X: adhwma_stat(X, cfg['lambda_min'], cfg['lambda_max'], cfg['gamma'])}
    rng = np.random.RandomState(cfg['seed'])
    X0 = rng.standard_normal((cfg['n_cal'], cfg['max_rl'], len(ctx.vars))).astype(np.float32)
    ucl, cal = {}, []
    for ch, f in fns.items():
        ctx.tick(f'Kalibrasi UCL {ch}', 0.3 if ch == 'MEWMA' else 0.6)
        s0 = f(X0); ucl[ch] = find_ucl(s0, 1 / cfg['alpha'])
        cal.append({'Chart': ch, 'UCL': ucl[ch], 'ARL0_target': 1 / cfg['alpha'],
                    'ARL0_hasil_kalibrasi': arl_hat(s0, ucl[ch])})
    rows, series = [], {}
    for (m_, w_) in best_keys:
        Z = standardize(ctx.Y_TEST - ctx.preds[(m_, w_)])[None].astype(np.float32)
        for ch, f in fns.items():
            s = f(Z)[0]; sig = s > ucl[ch]
            series[(m_, ch)] = (s, sig)
            rows.append({'Metode': m_, 'Window': w_, 'Chart': ch, 'UCL': ucl[ch], 'Total_Observasi': len(s),
                         'Out_of_Control': int(sig.sum()), 'Persen_OOC': 100 * sig.mean(),
                         'Episode_Sinyal': int(np.sum(np.diff(np.r_[0, sig.astype(int)]) == 1)),
                         'Sinyal_Pertama': str(pd.Timestamp(ctx.T_TEST[np.argmax(sig)])) if sig.any() else '-'})
    return pd.DataFrame(rows), pd.DataFrame(cal), series, ucl


# ----------------------------------------------------------------------------- orkestrasi
def run_analysis(df, cfg, progress=None):
    t_start = time.time()
    progress = progress or (lambda m, f: None)
    ctx = Ctx(df, cfg, progress); set_seed(cfg['seed'])
    methods = [m for m in METHOD_ORDER if cfg['run'].get(m)]
    n = len(methods)
    for i, m in enumerate(methods):
        ctx.base, ctx.span = i / (n + 1), 1 / (n + 1)
        progress(f'Melatih {m} …', ctx.base)
        t0 = time.time()
        try:
            RUNNERS[m](ctx)
        except ImportError as e:
            ctx.warnings.append(f'{m} dilewati karena library tidak tersedia: {e}')
        ctx.timing[m] = time.time() - t0
    if not ctx.results:
        raise ValueError('Tidak ada metode yang berhasil dijalankan.')

    ctx.base, ctx.span = n / (n + 1), 1 / (n + 1)
    progress('Menyusun peringkat …', ctx.base)
    sel = cfg['select_by']
    ALL = pd.DataFrame(ctx.results)
    BEST = ALL.loc[ALL.groupby('Metode')[sel].idxmin()].copy()
    BEST['Peringkat_RMSE'] = BEST['RMSE_avg'].rank(method='min').astype(int)
    BEST['Peringkat_MAPE'] = BEST['MAPE_avg'].rank(method='min').astype(int)
    BEST['Peringkat_Gabungan'] = ((BEST.Peringkat_RMSE + BEST.Peringkat_MAPE) / 2).rank(method='min').astype(int)
    BEST['Skor'] = 100 * BEST['RMSE_avg'].min() / BEST['RMSE_avg']
    BEST['Waktu_s'] = BEST['Metode'].map(ctx.timing)
    BEST = BEST.sort_values(['Peringkat_Gabungan', sel]).reset_index(drop=True)
    BEST.insert(0, 'No', range(1, len(BEST) + 1))
    best_keys = [(m, int(w)) for m, w in zip(BEST.Metode, BEST.Window)]

    win_all = ALL[ALL.Window.isin(cfg['windows'])]
    order = BEST.Metode.tolist()
    piv_r = win_all.pivot_table(index='Metode', columns='Window', values='RMSE_avg').reindex(order)
    piv_m = win_all.pivot_table(index='Metode', columns='Window', values='MAPE_avg').reindex(order)
    rows = []
    for v in ctx.vars:
        r1, r2 = BEST.loc[BEST[f'RMSE_{v}'].idxmin()], BEST.loc[BEST[f'MAPE_{v}'].idxmin()]
        rows.append({'Variabel': v, 'Terbaik_RMSE': f'{r1.Metode} (w={r1.Window})', 'RMSE': r1[f'RMSE_{v}'],
                     'Terbaik_MAPE': f'{r2.Metode} (w={r2.Window})', 'MAPE (%)': r2[f'MAPE_{v}']})
    PER_VAR = pd.DataFrame(rows)

    CC, CAL, CC_SERIES, UCL = run_control_charts(ctx, best_keys)
    if 'DL_RESULTS' in ctx.extra: ctx.extra['DL_RESULTS'] = pd.DataFrame(ctx.extra['DL_RESULTS'])
    return dict(cfg=cfg, ALL=ALL, BEST=BEST, PREDS=ctx.preds, Y_TEST=ctx.Y_TEST, T_TEST=ctx.T_TEST,
                SD_TRAIN=ctx.SD_TRAIN, N=ctx.N, N_TRAIN=ctx.N_TRAIN, VARS=ctx.vars, PIV_RMSE=piv_r, PIV_MAPE=piv_m,
                PER_VAR=PER_VAR, CC=CC, CAL=CAL, CC_SERIES=CC_SERIES, UCL=UCL, BEST_KEYS=best_keys,
                TIMING=ctx.timing, extra=ctx.extra, warnings=ctx.warnings, elapsed=time.time() - t_start,
                HP_TABLE=hp_table())


# ----------------------------------------------------------------------------- cache hasil di disk
def make_key(df, cfg):
    h = hashlib.md5()
    h.update(pd.util.hash_pandas_object(df, index=True).values.tobytes())
    h.update(json.dumps(cfg, sort_keys=True, default=str).encode())
    return h.hexdigest()[:16]


def save_cache(key, res):
    try:
        os.makedirs(CACHE_DIR, exist_ok=True)
        with open(os.path.join(CACHE_DIR, f'{key}.pkl'), 'wb') as f: pickle.dump(res, f)
    except Exception:
        pass


def load_cache(key):
    p = os.path.join(CACHE_DIR, f'{key}.pkl')
    if os.path.exists(p):
        try:
            with open(p, 'rb') as f: return pickle.load(f)
        except Exception:
            return None
    return None


# ----------------------------------------------------------------------------- ekspor Excel
def build_excel(res, data_name='dataset'):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    cfg, VARS, BEST, ALL = res['cfg'], res['VARS'], res['BEST'], res['ALL']
    WINDOWS, ex = cfg['windows'], res['extra']
    HEAD = PatternFill('solid', fgColor='1F4E78'); GOOD = PatternFill('solid', fgColor='63BE7B')
    ROWBEST = PatternFill('solid', fgColor='E2EFDA'); THIN = Side(style='thin', color='BFBFBF')

    def style_sheet(ws, df_, best_cols=(), best_row=None, title=None, n_dec=4):
        start = 1
        if title:
            ws.insert_rows(1, 2); ws['A1'] = title; ws['A1'].font = Font(bold=True, size=13); start = 3
        for c in range(1, df_.shape[1] + 1):
            cell = ws.cell(start, c); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = HEAD
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r in range(start + 1, start + 1 + len(df_)):
            for c in range(1, df_.shape[1] + 1):
                cell = ws.cell(r, c); cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
                if isinstance(cell.value, float): cell.number_format = '0.' + '0' * n_dec
                if best_row is not None and r - start - 1 == best_row: cell.fill = ROWBEST
        for col in best_cols:
            if col not in df_.columns: continue
            c = df_.columns.get_loc(col) + 1; vals = pd.to_numeric(df_[col], errors='coerce')
            if vals.notna().any():
                r = int(np.nanargmin(vals.values)) + start + 1
                ws.cell(r, c).fill = GOOD; ws.cell(r, c).font = Font(bold=True)
        for c in range(1, df_.shape[1] + 1):
            L = max([len(str(df_.columns[c - 1]))] + [len(str(x)) for x in df_.iloc[:, c - 1].astype(str).head(200)])
            ws.column_dimensions[get_column_letter(c)].width = min(max(L + 2, 10), 60)
        ws.freeze_panes = ws.cell(start + 1, 1)

    def write(wr, df_, sheet, **kw):
        df_.to_excel(wr, sheet_name=sheet, index=False, startrow=0)
        style_sheet(wr.book[sheet], df_, **kw)

    ring = BEST[['No', 'Metode', 'Kategori', 'Window', 'RMSE_avg', 'MAPE_avg', 'RMSE_norm', 'Peringkat_RMSE',
                 'Peringkat_MAPE', 'Peringkat_Gabungan', 'Konfigurasi', 'Catatan']].copy()
    ring.columns = ['No', 'Metode', 'Kategori', 'Window Terbaik', 'RMSE Rata-rata', 'MAPE Rata-rata (%)',
                    'RMSE Ternormalisasi', 'Peringkat RMSE', 'Peringkat MAPE', 'Peringkat Gabungan', 'Konfigurasi', 'Catatan']

    def per_var(kind):
        t = BEST[['Metode', 'Window'] + [f'{kind}_{v}' for v in VARS] + [f'{kind}_avg']].copy()
        t.columns = ['Metode', 'Window'] + VARS + ['Rata-rata']; return t

    def piv_out(p):
        t = p.copy(); t.columns = [f'Window {c}' for c in t.columns]; t = t.reset_index()
        t['Terbaik'] = p.min(1).values; return t

    MCOLS = [f'RMSE_{v}' for v in VARS] + [f'MAPE_{v}' for v in VARS] + ['RMSE_avg', 'MAPE_avg', 'RMSE_norm']
    semua = ALL.sort_values(['Metode', 'Window'])[['Metode', 'Kategori', 'Window'] + MCOLS + ['Konfigurasi', 'Catatan']]
    pred_out = pd.DataFrame({'Time': pd.to_datetime(res['T_TEST'])})
    for j, v in enumerate(VARS): pred_out[f'{v}_Aktual'] = res['Y_TEST'][:, j]
    for (m_, w_) in res['BEST_KEYS']:
        for j, v in enumerate(VARS): pred_out[f'{v}_{m_}_w{w_}'] = res['PREDS'][(m_, w_)][:, j]
    top = BEST.iloc[0]
    info = pd.DataFrame({'Keterangan': ['Data', 'Split', 'Prediksi', 'Kriteria window terbaik', 'Peringkat Gabungan',
                                        'Metode terbaik (gabungan)', 'RMSE terkecil', 'MAPE terkecil', 'Waktu komputasi (detik)'],
                         'Isi': [f"{data_name} ({res['N']} observasi)",
                                 f"{cfg['train_ratio']*100:.0f}% train ({res['N_TRAIN']}) : test ({res['N'] - res['N_TRAIN']}), kronologis",
                                 f"One-step-ahead multivariat, sliding window {WINDOWS}", f"{cfg['select_by']} terkecil",
                                 'Rata-rata peringkat RMSE dan MAPE (1 = terbaik)', f'{top.Metode} (window {top.Window})',
                                 f"{BEST.loc[BEST.RMSE_avg.idxmin(), 'Metode']} = {BEST.RMSE_avg.min():.4f}",
                                 f"{BEST.loc[BEST.MAPE_avg.idxmin(), 'Metode']} = {BEST.MAPE_avg.min():.3f}%",
                                 ', '.join(f'{k}: {v:.0f}' for k, v in res['TIMING'].items())]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as wr:
        write(wr, ring, '1_Ringkasan_Terbaik', best_cols=['RMSE Rata-rata', 'MAPE Rata-rata (%)', 'RMSE Ternormalisasi'],
              best_row=0, title='Perbandingan metode (window terbaik tiap metode)')
        write(wr, per_var('RMSE'), '2_RMSE_per_Variabel', best_cols=VARS + ['Rata-rata'], title='RMSE per variabel')
        write(wr, per_var('MAPE'), '3_MAPE_per_Variabel', best_cols=VARS + ['Rata-rata'], title='MAPE (%) per variabel')
        write(wr, piv_out(res['PIV_RMSE']), '4_RMSE_per_Window', best_cols=[f'Window {w}' for w in WINDOWS] + ['Terbaik'], title='RMSE rata-rata: metode x window')
        write(wr, piv_out(res['PIV_MAPE']), '5_MAPE_per_Window', best_cols=[f'Window {w}' for w in WINDOWS] + ['Terbaik'], title='MAPE rata-rata (%): metode x window')
        write(wr, res['PER_VAR'], '6_Terbaik_per_Variabel', title='Metode terbaik untuk setiap variabel')
        write(wr, res['CC'], '7_Control_Chart', title='Control chart residual data uji')
        write(wr, res['CAL'], '8_Kalibrasi_UCL', title='Kalibrasi UCL Monte Carlo')
        write(wr, semua, 'Lampiran_Semua_Hasil', best_cols=['RMSE_avg', 'MAPE_avg'])
        write(wr, res['HP_TABLE'], 'Lampiran_HP_XGB_LGBM', n_dec=6)
        if 'SVR_BEST' in ex: write(wr, ex['SVR_BEST'], 'Lampiran_SVR_Terbaik'); write(wr, ex['SVR_TUNING'], 'Lampiran_SVR_Tuning')
        if 'RF_BEST' in ex: write(wr, ex['RF_BEST'], 'Lampiran_RF_Terbaik')
        if 'DL_RESULTS' in ex and len(ex['DL_RESULTS']): write(wr, ex['DL_RESULTS'], 'Lampiran_DL_Grid', best_cols=['RMSE_avg', 'MAPE_avg'])
        if 'VAR_ADF' in ex:
            write(wr, ex['VAR_ADF'], 'Lampiran_VAR_ADF'); write(wr, ex['VAR_LAG'], 'Lampiran_VAR_Lag'); write(wr, ex['VAR_STAB'], 'Lampiran_VAR_Stabilitas')
        write(wr, pred_out, 'Lampiran_Prediksi_Uji')
        write(wr, info, 'Info')
    return buf.getvalue()
